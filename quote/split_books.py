# -*- coding: utf-8 -*-
"""提出用の3点を1ファイルずつの Excel に切り出す

  python3 quote/split_books.py

原本（全7シート）から 見積書 / 注文書 / 前提条件 をそれぞれ単独ブックにします。
他シートを参照している数式は計算結果の値に置き換えるため、
切り出したファイルは単独で開いても数字が壊れません。
数字を変えるときは原本を直してから、このスクリプトを流し直してください。
"""
import os, shutil, subprocess, tempfile
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
BOOK = os.path.join(HERE, "聖建様_概算見積書・注文書.xlsx")

TARGETS = [
    ("01_見積書", "聖建様_見積書.xlsx"),
    ("06_注文書", "聖建様_注文書.xlsx"),
    ("07_前提条件", "聖建様_前提条件.xlsx"),
]


def recalculated_values():
    """LibreOffice に再計算させ、全セルの計算結果を取り出す"""
    tmp = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx",
                    "--outdir", tmp, BOOK], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    path = os.path.join(tmp, os.path.basename(BOOK))
    wb = openpyxl.load_workbook(path, data_only=True)
    vals = {ws.title: {c.coordinate: c.value for r in ws.iter_rows() for c in r}
            for ws in wb}
    wb.close()
    shutil.rmtree(tmp, ignore_errors=True)
    return vals


def main():
    vals = recalculated_values()

    for keep, outname in TARGETS:
        wb = openpyxl.load_workbook(BOOK)

        # 数式を計算結果に置き換える（他シートへの参照を断ち切る）
        ws = wb[keep]
        n = 0
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = vals[keep].get(c.coordinate)
                    n += 1

        for name in list(wb.sheetnames):
            if name != keep:
                del wb[name]

        out = os.path.join(HERE, outname)
        wb.save(out)
        print("書き出し: %s（数式 %d 件を値に置換）" % (outname, n))


if __name__ == "__main__":
    main()
